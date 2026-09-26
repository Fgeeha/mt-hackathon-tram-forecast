"""Геопривязка: маршруты и остановки → artifacts/geo.json.

Источники: справочник организаторов (лист «Порядок_с_координатами», маршруты
1, 5, 7, 11, 12) и OpenStreetMap (relation route=tram) для остальных маршрутов
и для геометрии путей. Валидации не содержат остановки, поэтому прогноз по
остановке — это прогноз маршрута × вес остановки (гравитационная оценка:
пересадочные узлы и конечные притягивают больше пассажиров).
"""

import json
import logging
import re

import openpyxl

from tram_forecast.settings import ARTIFACTS, EXTERNAL, RAW, ROUTES

log = logging.getLogger(__name__)

HUB = re.compile(r"метро|мцк|мцд|вокзал|станци|платформ|рынок|больниц|парк", re.IGNORECASE)
SPRAV = RAW / "spravochniki" / "Хакатон_справочники_трамвай_10_маршрутов.xlsx"


def stop_weight(name: str, terminal: bool) -> float:
    return 1.0 + 2.0 * bool(HUB.search(name or "")) + 1.0 * terminal


def from_sprav() -> dict[int, dict]:
    ws = openpyxl.load_workbook(SPRAV, read_only=True)["Порядок_с_координатами"]
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    recs = [dict(zip(head, r)) for r in rows[1:] if r[0]]
    names = {r[2]: r[3] for r in list(openpyxl.load_workbook(SPRAV, read_only=True)["Маршруты GTFS_ROUTES"].iter_rows(values_only=True))[2:]}
    out: dict[int, dict] = {}
    for rec in recs:
        route = int(rec["route_short_name"])
        if route not in ROUTES:
            continue
        r = out.setdefault(route, {"name": names.get(str(route), ""), "source": "справочник", "dirs": {}})
        r["dirs"].setdefault(int(rec["direction_id"]), []).append(
            (int(rec["stop_sequence"]), str(rec["stop_id"]), rec["stop_name"], float(rec["stop_lat"]), float(rec["stop_lon"]))
        )
    for r in out.values():
        r["lines"], r["stops"] = [], []
        for d, seq in sorted(r.pop("dirs").items()):
            seq.sort()
            r["lines"].append([[lon, lat] for _, _, _, lat, lon in seq])
            for i, (s, sid, name, lat, lon) in enumerate(seq):
                r["stops"].append({"stop_id": sid, "name": name, "lat": lat, "lon": lon, "dir": d, "seq": s,
                                   "weight": stop_weight(name, i in (0, len(seq) - 1))})
    return out


def from_osm() -> dict[int, dict]:
    path = EXTERNAL / "osm_tram_routes.json"
    if not path.exists():
        return {}
    out: dict[int, dict] = {}
    for rel in json.loads(path.read_text())["elements"]:
        tags = rel.get("tags", {})
        if not tags.get("ref", "").isdigit() or int(tags["ref"]) not in ROUTES:
            continue
        route = int(tags["ref"])
        r = out.setdefault(route, {"name": tags.get("name", ""), "source": "OpenStreetMap", "lines": [], "stops": []})
        d = len({s["dir"] for s in r["stops"]})  # каждое направление — отдельная relation
        stops = [m for m in rel.get("members", []) if m["type"] == "node" and m.get("role", "").startswith("stop")]
        for i, m in enumerate(stops):
            r["stops"].append({"stop_id": f"osm{m['ref']}", "name": f"Остановка {i + 1}", "lat": m["lat"], "lon": m["lon"],
                               "dir": d, "seq": i + 1, "weight": stop_weight("", i in (0, len(stops) - 1))})
        for m in rel.get("members", []):
            if m["type"] == "way" and m.get("geometry"):
                r["lines"].append([[p["lon"], p["lat"]] for p in m["geometry"]])
    return out


def build_geo() -> dict:
    sprav, osm = from_sprav(), from_osm()
    geo = {}
    for route in ROUTES:
        g = sprav.get(route) or osm.get(route)
        if not g:
            continue
        if route in sprav and route in osm:
            g["lines"] = osm[route]["lines"]  # точная геометрия путей из OSM
        total = sum(s["weight"] for s in g["stops"]) or 1.0
        for s in g["stops"]:
            s["share"] = round(s["weight"] / total, 5)
        geo[str(route)] = g
    return geo


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    geo = build_geo()
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    (ARTIFACTS / "geo.json").write_text(json.dumps(geo, ensure_ascii=False))
    log.info("geo: %s", {k: (v["source"], len(v["stops"])) for k, v in geo.items()})


if __name__ == "__main__":
    main()
